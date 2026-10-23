"""
utils.py
Helper functions for the AI Resume & Cover Letter Generator (powered by Google Gemini).

Responsibilities:
- Extract raw text from an uploaded resume (PDF / DOCX / TXT)
- Call the Gemini API (google-genai SDK) to generate a tailored resume + cover letter
- Build downloadable PDF files from the (possibly user-edited) text
"""

from pypdf import PdfReader
import docx
from fpdf import FPDF
from google.genai import types


# ---------------------------------------------------------------------------
# 1. Extract text from the uploaded resume file
# ---------------------------------------------------------------------------
def extract_text_from_file(uploaded_file) -> str:
    """Read a Streamlit UploadedFile (pdf, docx, or txt) and return plain text."""
    name = uploaded_file.name.lower()
    uploaded_file.seek(0)

    if name.endswith(".pdf"):
        reader = PdfReader(uploaded_file)
        text = "\n".join((page.extract_text() or "") for page in reader.pages)

    elif name.endswith(".docx"):
        document = docx.Document(uploaded_file)
        text = "\n".join(p.text for p in document.paragraphs)

    else:  # assume plain text
        text = uploaded_file.read().decode("utf-8", errors="ignore")

    return text.strip()


# ---------------------------------------------------------------------------
# 2. Call Gemini to generate a tailored resume + cover letter
# ---------------------------------------------------------------------------
RESUME_MARKER = "===RESUME==="
COVER_MARKER = "===COVER_LETTER==="

SYSTEM_PROMPT = """You are an expert resume writer and career coach.
You will be given a candidate's existing resume (raw text) and a job description.

Produce two things, staying strictly truthful to the candidate's real experience
(never invent employers, titles, dates, or degrees):

1. A TAILORED RESUME: reorganize, reword, and emphasize the candidate's real
   skills/experience so it aligns with the job description, naturally including
   relevant keywords for ATS screening. Keep it well-structured and ready to
   read as a plain-text resume (clear section headers, bullet points using "- ").

2. A TAILORED COVER LETTER: a concise, personalized cover letter (3-4 short
   paragraphs) connecting the candidate's real background to this specific role
   and company/industry context implied by the job description.

Output EXACTLY in this format, with nothing before, after, or outside it:

===RESUME===
<full tailored resume text here>
===COVER_LETTER===
<full cover letter text here>
"""


def generate_tailored_content(client, resume_text: str, job_description: str, model: str = "gemini-2.5-flash"):
    """Call the Gemini API and return (tailored_resume, cover_letter) as strings."""
    user_prompt = (
        f"CANDIDATE'S CURRENT RESUME:\n{resume_text}\n\n"
        f"JOB DESCRIPTION:\n{job_description}\n"
    )

    response = client.models.generate_content(
        model=model,
        contents=user_prompt,
        config=types.GenerateContentConfig(
            system_instruction=SYSTEM_PROMPT,
            max_output_tokens=4000,
        ),
    )

    full_text = response.text or ""
    return _split_output(full_text)


def _split_output(full_text: str):
    """Split the model's raw output into (resume, cover_letter) using the markers."""
    resume_part, cover_part = "", ""

    if RESUME_MARKER in full_text and COVER_MARKER in full_text:
        after_resume_marker = full_text.split(RESUME_MARKER, 1)[1]
        resume_part, cover_part = after_resume_marker.split(COVER_MARKER, 1)
    else:
        # Fallback: if the model didn't follow the format, put everything in the resume box
        resume_part = full_text
        cover_part = ""

    return resume_part.strip(), cover_part.strip()


# ---------------------------------------------------------------------------
# 3. Build a downloadable PDF from text
# ---------------------------------------------------------------------------
def create_pdf(text: str, title: str = "Document") -> bytes:
    """Turn plain text into a simple, clean single/multi-page PDF and return its bytes."""
    pdf = FPDF(format="A4", unit="mm")
    pdf.set_auto_page_break(auto=True, margin=18)
    pdf.add_page()
    pdf.set_margins(18, 18, 18)

    pdf.set_font("Helvetica", "B", 15)
    pdf.cell(0, 10, title, ln=True)
    pdf.ln(2)
    pdf.set_font("Helvetica", size=11)

    for raw_line in text.split("\n"):
        line = raw_line.rstrip()
        if not line:
            pdf.ln(4)
            continue
        # Encode to latin-1 (PDF core fonts only support this range);
        # replace anything unsupported so generation never crashes.
        safe_line = line.encode("latin-1", errors="replace").decode("latin-1")
        # wrapmode="CHAR" lets fpdf2 break a line at ANY character when needed
        # (not just at spaces), so a long URL/email/unbroken token can never
        # cause "Not enough horizontal space to render a single character".
        try:
            pdf.multi_cell(0, 6, safe_line, wrapmode="CHAR")
        except Exception:
            # Extremely defensive fallback: strip anything unusual and retry
            # once, so PDF generation can never crash the app.
            fallback_line = "".join(
                ch if ch.isalnum() or ch in " .,:;()-/@_" else " " for ch in safe_line
            ).strip()
            if fallback_line:
                pdf.multi_cell(0, 6, fallback_line, wrapmode="CHAR")

    output = pdf.output(dest="S")
    if isinstance(output, str):
        output = output.encode("latin-1", errors="replace")
    return bytes(output)
